from __future__ import annotations

import asyncio
from collections import defaultdict
from datetime import date, datetime, timezone
from io import BytesIO
import json
import re
from urllib import request
from urllib.error import URLError
from urllib.parse import quote

from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from app.core.config import get_settings
from app.schemas.expense import ExpenseCategorySummary, ExpenseImportResponse, ExpenseTransaction


MAX_EXTRACTED_CHARS = 120_000
SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}

CATEGORIES = [
    "Food & Dining",
    "Groceries",
    "Shopping",
    "Transport",
    "Travel",
    "Bills & Utilities",
    "Rent & Housing",
    "Healthcare",
    "Entertainment",
    "Education",
    "Fees & Charges",
    "Transfers",
    "Investments",
    "Income",
    "Other",
]

CATEGORY_KEYWORDS = {
    "Food & Dining": ["restaurant", "cafe", "swiggy", "zomato", "food", "dining"],
    "Groceries": ["grocery", "supermarket", "bigbasket", "blinkit", "zepto", "dmart"],
    "Shopping": ["amazon", "flipkart", "myntra", "shopping", "store", "retail"],
    "Transport": ["uber", "ola", "metro", "fuel", "petrol", "diesel", "transport"],
    "Travel": ["flight", "hotel", "airbnb", "irctc", "makemytrip", "goibibo"],
    "Bills & Utilities": ["bill", "electricity", "water", "gas", "broadband", "mobile", "recharge"],
    "Rent & Housing": ["rent", "maintenance", "housing"],
    "Healthcare": ["hospital", "pharmacy", "doctor", "medical", "health"],
    "Entertainment": ["netflix", "spotify", "prime", "movie", "bookmyshow", "entertainment"],
    "Education": ["school", "college", "course", "tuition", "education"],
    "Fees & Charges": ["fee", "charge", "gst", "penalty", "atm"],
    "Transfers": ["transfer", "upi", "neft", "imps", "rtgs"],
    "Investments": ["mutual fund", "zerodha", "groww", "investment", "sip"],
    "Income": ["salary", "interest", "cashback", "refund", "dividend"],
}


class ExpenseTrackerError(Exception):
    pass


def extract_statement_text(filename: str, content_type: str | None, file_bytes: bytes) -> str:
    lowered_name = filename.lower()

    if content_type == "application/pdf" or lowered_name.endswith(".pdf"):
        return _extract_pdf_text(file_bytes)

    if (
        content_type in SUPPORTED_CONTENT_TYPES
        or lowered_name.endswith(".xlsx")
        or lowered_name.endswith(".xls")
    ):
        return _extract_excel_text(file_bytes)

    raise ExpenseTrackerError("Upload a PDF or Excel account statement.")


async def build_expense_report(filename: str, content_type: str | None, file_bytes: bytes) -> ExpenseImportResponse:
    statement_text = extract_statement_text(filename, content_type, file_bytes)
    if not statement_text.strip():
        raise ExpenseTrackerError("The statement did not contain readable transaction text.")

    settings = get_settings()
    parser = "heuristic"
    transactions: list[ExpenseTransaction]

    if settings.llm_provider == "bedrock" and settings.bedrock_api_key:
        try:
            transactions = await _parse_with_bedrock(
                statement_text,
                settings.bedrock_api_key,
                settings.bedrock_region,
                settings.bedrock_model,
            )
            parser = "bedrock"
        except (ExpenseTrackerError, ValidationError, json.JSONDecodeError, URLError):
            transactions = _parse_with_heuristics(statement_text)
    elif settings.openai_api_key:
        try:
            transactions = await _parse_with_llm(statement_text, settings.openai_api_key, settings.openai_model)
            parser = "openai"
        except (ExpenseTrackerError, ValidationError, json.JSONDecodeError, URLError):
            transactions = _parse_with_heuristics(statement_text)
    else:
        transactions = _parse_with_heuristics(statement_text)

    transactions = _dedupe_transactions(transactions)
    if not transactions:
        raise ExpenseTrackerError("No expense transactions could be identified from this statement.")

    return _build_response(filename, parser, transactions)


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)[:MAX_EXTRACTED_CHARS]


def _extract_excel_text(file_bytes: bytes) -> str:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    lines: list[str] = []

    for worksheet in workbook.worksheets:
        lines.append(f"Sheet: {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                lines.append("\t".join(values))

    return "\n".join(lines)[:MAX_EXTRACTED_CHARS]


async def _parse_with_llm(statement_text: str, api_key: str, model: str) -> list[ExpenseTransaction]:
    raw_response = await asyncio.to_thread(_call_openai_responses_api, statement_text, api_key, model)
    parsed = _extract_response_json(raw_response)
    return [ExpenseTransaction.model_validate(item) for item in parsed.get("transactions", [])]


async def _parse_with_bedrock(
    statement_text: str,
    api_key: str,
    region: str,
    model: str,
) -> list[ExpenseTransaction]:
    raw_response = await asyncio.to_thread(_call_bedrock_converse_api, statement_text, api_key, region, model)
    parsed = _extract_bedrock_json(raw_response)
    return [ExpenseTransaction.model_validate(item) for item in parsed.get("transactions", [])]


def _expense_prompt(statement_text: str) -> str:
    return (
        "Parse this account statement into normalized expense transactions.\n"
        "Return only valid JSON with this exact shape: "
        '{"transactions":[{"transaction_date":"YYYY-MM-DD","description":"string",'
        '"merchant":"string or null","category":"one allowed category","amount":123.45,'
        '"currency":"INR","account_name":"string or null","confidence":0.0}]}\n'
        "Rules:\n"
        "- Extract only outgoing personal expenses from bank or card statements.\n"
        "- Ignore opening/closing balances, failed transactions, and pure account metadata.\n"
        "- Use positive amounts for spends.\n"
        "- Treat refunds, salary, interest, and investment redemptions as Income only when clearly credits.\n"
        "- Return INR unless another currency is explicit.\n"
        f"- Categories must be one of: {', '.join(CATEGORIES)}.\n\n"
        f"Statement text:\n{statement_text}"
    )


def _call_bedrock_converse_api(statement_text: str, api_key: str, region: str, model: str) -> dict:
    payload = {
        "messages": [
            {
                "role": "user",
                "content": [{"text": _expense_prompt(statement_text)}],
            }
        ],
        "inferenceConfig": {
            "maxTokens": 12000,
            "temperature": 0.0,
        },
    }

    encoded_model = quote(model, safe="")
    api_request = request.Request(
        f"https://bedrock-runtime.{region}.amazonaws.com/model/{encoded_model}/converse",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    with request.urlopen(api_request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def _call_openai_responses_api(statement_text: str, api_key: str, model: str) -> dict:
    payload = {
        "model": model,
        "instructions": (
            "You are an expense extraction engine. Extract only outgoing personal expenses from "
            "bank or card statements. Ignore opening/closing balances, failed transactions, and pure "
            "account metadata. Treat refunds, salary, interest, and investment redemptions as Income "
            "only when they are clearly credits. Return INR unless another currency is explicit."
        ),
        "input": (
            "Parse this account statement into normalized expense transactions. "
            "Use positive amounts for spends. Categories must be one of: "
            f"{', '.join(CATEGORIES)}.\n\nStatement text:\n{statement_text}"
        ),
        "text": {
            "format": {
                "type": "json_schema",
                "name": "expense_transactions",
                "strict": True,
                "schema": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "transactions": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "properties": {
                                    "transaction_date": {"type": "string"},
                                    "description": {"type": "string"},
                                    "merchant": {"type": ["string", "null"]},
                                    "category": {"type": "string", "enum": CATEGORIES},
                                    "amount": {"type": "number"},
                                    "currency": {"type": "string"},
                                    "account_name": {"type": ["string", "null"]},
                                    "confidence": {"type": "number"},
                                },
                                "required": [
                                    "transaction_date",
                                    "description",
                                    "merchant",
                                    "category",
                                    "amount",
                                    "currency",
                                    "account_name",
                                    "confidence",
                                ],
                            },
                        }
                    },
                    "required": ["transactions"],
                },
            }
        },
    }

    api_request = request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    with request.urlopen(api_request, timeout=45) as response:
        return json.loads(response.read().decode("utf-8"))


def _extract_response_json(response_body: dict) -> dict:
    for output in response_body.get("output", []):
        for content in output.get("content", []):
            if content.get("type") == "output_text":
                return json.loads(content.get("text", "{}"))

    raise ExpenseTrackerError("The LLM did not return structured expense data.")


def _extract_bedrock_json(response_body: dict) -> dict:
    text_chunks = []
    for content in response_body.get("output", {}).get("message", {}).get("content", []):
        if "text" in content:
            text_chunks.append(content["text"])

    raw_text = "\n".join(text_chunks).strip()
    if not raw_text:
        raise ExpenseTrackerError("Bedrock did not return expense data.")

    return json.loads(_extract_json_object(raw_text))


def _extract_json_object(raw_text: str) -> str:
    cleaned = raw_text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)

    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ExpenseTrackerError("The LLM response did not include a JSON object.")

    return cleaned[start : end + 1]


def _parse_with_heuristics(statement_text: str) -> list[ExpenseTransaction]:
    transactions: list[ExpenseTransaction] = []
    date_pattern = re.compile(r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|\d{4}-\d{1,2}-\d{1,2})\b")
    amount_pattern = re.compile(r"(?:INR|Rs\.?|₹)?\s*([0-9][0-9,]*(?:\.\d{1,2})?)")

    for raw_line in statement_text.splitlines():
        line = " ".join(raw_line.split())
        if not line:
            continue

        date_match = date_pattern.search(line)
        amount_matches = amount_pattern.findall(line)
        if not date_match or not amount_matches:
            continue

        amount = _parse_amount(amount_matches[-1])
        if amount <= 0:
            continue

        description = line.replace(date_match.group(0), "").strip(" -|\t")
        if not description:
            description = line

        if _looks_like_credit(line):
            continue

        transactions.append(
            ExpenseTransaction(
                transaction_date=_parse_date(date_match.group(0)),
                description=description[:300],
                merchant=_guess_merchant(description),
                category=_categorize(description),
                amount=amount,
                currency="INR",
                account_name=None,
                confidence=0.45,
            )
        )

    return transactions


def _parse_amount(value: str) -> float:
    return round(float(value.replace(",", "")), 2)


def _parse_date(value: str) -> date:
    for pattern in ("%d-%m-%Y", "%d/%m/%Y", "%d-%m-%y", "%d/%m/%y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            continue

    raise ExpenseTrackerError(f"Unable to parse transaction date: {value}")


def _looks_like_credit(line: str) -> bool:
    lowered = line.lower()
    credit_markers = [" credit", " cr", " deposit", " salary", " refund", " interest"]
    debit_markers = [" debit", " dr", " withdrawal", " purchase", " paid"]

    return any(marker in lowered for marker in credit_markers) and not any(
        marker in lowered for marker in debit_markers
    )


def _categorize(description: str) -> str:
    lowered = description.lower()
    for category, keywords in CATEGORY_KEYWORDS.items():
        if any(keyword in lowered for keyword in keywords):
            return category

    return "Other"


def _guess_merchant(description: str) -> str | None:
    cleaned = re.sub(r"\b(?:upi|pos|neft|imps|rtgs|debit|card|txn|ref|id)\b", " ", description, flags=re.I)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -*/")
    if not cleaned:
        return None

    return cleaned[:120]


def _dedupe_transactions(transactions: list[ExpenseTransaction]) -> list[ExpenseTransaction]:
    seen: set[tuple[date, str, float]] = set()
    deduped: list[ExpenseTransaction] = []

    for transaction in transactions:
        key = (
            transaction.transaction_date,
            transaction.description.lower().strip(),
            round(transaction.amount, 2),
        )
        if key in seen:
            continue
        seen.add(key)
        deduped.append(transaction)

    return sorted(deduped, key=lambda item: (item.transaction_date, item.description))


def _build_response(filename: str, parser: str, transactions: list[ExpenseTransaction]) -> ExpenseImportResponse:
    totals_by_category: dict[str, float] = defaultdict(float)
    counts_by_category: dict[str, int] = defaultdict(int)

    for transaction in transactions:
        totals_by_category[transaction.category] += transaction.amount
        counts_by_category[transaction.category] += 1

    category_totals = [
        ExpenseCategorySummary(
            category=category,
            total=round(total, 2),
            transaction_count=counts_by_category[category],
        )
        for category, total in sorted(totals_by_category.items(), key=lambda item: item[1], reverse=True)
    ]

    return ExpenseImportResponse(
        id="",
        source_filename=filename,
        parser=parser,
        statement_period_start=min(transaction.transaction_date for transaction in transactions),
        statement_period_end=max(transaction.transaction_date for transaction in transactions),
        total_spend=round(sum(transaction.amount for transaction in transactions), 2),
        currency=transactions[0].currency if transactions else "INR",
        category_totals=category_totals,
        transactions=transactions,
        created_at=datetime.now(timezone.utc),
    )
